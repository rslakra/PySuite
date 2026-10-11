"""Tests for ``pycrux.converter.markdown.Word2MarkdownConverter`` (uses the real ``python-docx``)."""

import pytest
from docx import Document

import pycrux.converter.markdown as markdown_module
from pycrux.converter import AbstractConverter
from pycrux.converter.markdown import Word2MarkdownConverter


def _make_docx(path, build) -> None:
    doc = Document()
    build(doc)
    doc.save(path)


def test_class_contract() -> None:
    assert issubclass(Word2MarkdownConverter, AbstractConverter)
    assert Word2MarkdownConverter.source_extensions == {"docx"}
    assert Word2MarkdownConverter.dest_extensions == {"md", "markdown"}


def test_converts_headings_lists_and_paragraphs(tmp_path) -> None:
    def build(doc) -> None:
        doc.add_heading("Title", level=1)
        doc.add_heading("Section", level=2)
        doc.add_heading("Sub", level=3)
        doc.add_paragraph("first bullet", style="List Bullet")
        doc.add_paragraph("first step", style="List Number")
        doc.add_paragraph("plain text")

    source, destination = tmp_path / "in.docx", tmp_path / "out.md"
    _make_docx(source, build)

    Word2MarkdownConverter().convert(source, destination)

    assert destination.read_text(encoding="utf-8") == "\n".join(
        ["# Title", "## Section", "### Sub", "- first bullet", "1. first step", "plain text"]
    )


def test_blank_and_whitespace_paragraphs_become_empty_lines(tmp_path) -> None:
    def build(doc) -> None:
        doc.add_paragraph("a")
        doc.add_paragraph("")
        doc.add_paragraph("   ")
        doc.add_paragraph("b")

    source, destination = tmp_path / "in.docx", tmp_path / "out.md"
    _make_docx(source, build)
    Word2MarkdownConverter().convert(source, destination)
    assert destination.read_text(encoding="utf-8") == "a\n\n\nb"


def test_title_style_is_not_treated_as_a_heading(tmp_path) -> None:
    source, destination = tmp_path / "in.docx", tmp_path / "out.md"
    _make_docx(source, lambda doc: doc.add_heading("Doc title", level=0))
    Word2MarkdownConverter().convert(source, destination)
    assert destination.read_text(encoding="utf-8") == "Doc title"


def test_text_is_stripped(tmp_path) -> None:
    source, destination = tmp_path / "in.docx", tmp_path / "out.md"
    _make_docx(source, lambda doc: doc.add_paragraph("  padded  "))
    Word2MarkdownConverter().convert(source, destination)
    assert destination.read_text(encoding="utf-8") == "padded"


def test_empty_document_creates_an_empty_file(tmp_path) -> None:
    source, destination = tmp_path / "in.docx", tmp_path / "out.md"
    _make_docx(source, lambda doc: None)
    Word2MarkdownConverter().convert(source, destination)
    assert destination.read_text(encoding="utf-8") == ""


def test_creates_missing_destination_directories_and_accepts_strings(tmp_path) -> None:
    source = tmp_path / "in.docx"
    _make_docx(source, lambda doc: doc.add_paragraph("x"))
    destination = tmp_path / "deep" / "er" / "out.md"
    Word2MarkdownConverter().convert(str(source), str(destination))  # type: ignore[arg-type]
    assert destination.read_text(encoding="utf-8") == "x"


def test_unicode_is_preserved(tmp_path) -> None:
    source, destination = tmp_path / "in.docx", tmp_path / "out.md"
    _make_docx(source, lambda doc: doc.add_paragraph("héllo ✓ 日本語"))
    Word2MarkdownConverter().convert(source, destination)
    assert destination.read_text(encoding="utf-8") == "héllo ✓ 日本語"


def test_missing_source_raises(tmp_path) -> None:
    with pytest.raises(Exception):
        Word2MarkdownConverter().convert(tmp_path / "missing.docx", tmp_path / "out.md")


def test_reports_a_missing_python_docx_library(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(markdown_module, "Document", None)
    with pytest.raises(ImportError, match="pip install python-docx"):
        Word2MarkdownConverter().convert(tmp_path / "in.docx", tmp_path / "out.md")


def test_module_imports_without_python_docx(reload_without, tmp_path) -> None:
    module = reload_without(markdown_module, "docx")
    assert module.Document is None
    with pytest.raises(ImportError, match="pip install python-docx"):
        module.Word2MarkdownConverter().convert(tmp_path / "in.docx", tmp_path / "out.md")
