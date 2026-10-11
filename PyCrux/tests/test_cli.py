"""Tests for ``pycrux.cli`` (document converter CLI) and the root ``cli.py`` shim."""

import argparse
import importlib.util
from pathlib import Path

import pytest
from docx import Document

import pycrux.cli as cli

# --- _normalize_argv ---------------------------------------------------------------


def test_normalize_argv() -> None:
    assert cli._normalize_argv([]) == []
    assert cli._normalize_argv(["convert", "-i", "a.md"]) == ["-i", "a.md"]
    assert cli._normalize_argv(["convert"]) == []
    assert cli._normalize_argv(["list"]) == ["--list"]
    assert cli._normalize_argv(["list", "extra"]) == ["--list"]
    assert cli._normalize_argv(["-i", "a.md"]) == ["-i", "a.md"]
    assert cli._normalize_argv(["a.md", "convert"]) == ["a.md", "convert"]


# --- _build_parser -----------------------------------------------------------------


def test_parser_name_and_examples(capsys) -> None:
    parser = cli._build_parser()
    assert parser.prog == "pycrux.cli"
    help_text = parser.format_help()
    assert "python -m pycrux.cli -i README.md" in help_text
    assert "core.cli" not in help_text


def test_parser_parses_every_option() -> None:
    args = cli._build_parser().parse_args(["-i", "a.md", "--out", "b.docx", "-f", "pdf"])
    assert args.input_path == Path("a.md")
    assert args.output_path == Path("b.docx")
    assert args.dest_format == "pdf"
    assert args.list is False and args.source is None and args.destination is None


def test_parser_hidden_aliases_set_dest_format() -> None:
    parser = cli._build_parser()
    assert parser.parse_args(["--output", "pdf"]).dest_format == "pdf"
    assert parser.parse_args(["--to", "md"]).dest_format == "md"
    assert "--output" not in parser.format_help() and "--to" not in parser.format_help()


def test_parser_positionals() -> None:
    args = cli._build_parser().parse_args(["a.md", "b.docx"])
    assert (args.source, args.destination) == (Path("a.md"), Path("b.docx"))


def test_parser_list_flag() -> None:
    assert cli._build_parser().parse_args(["--list"]).list is True


# --- _default_destination / _resolve_paths -----------------------------------------


def test_default_destination() -> None:
    assert cli._default_destination(Path("/x/a.md")) == Path("/x/a.docx")
    assert cli._default_destination(Path("/x/a.docx"), "pdf") == Path("/x/a.pdf")
    with pytest.raises(ValueError):
        cli._default_destination(Path("a.txt"))


def test_resolve_paths_prefers_flags_over_positionals() -> None:
    args = argparse.Namespace(
        input_path=Path("flag.md"), source=Path("pos.md"), output_path=Path("flag.docx"), destination=Path("pos.docx")
    )
    assert cli._resolve_paths(args) == (Path("flag.md"), Path("flag.docx"))
    args = argparse.Namespace(input_path=None, source=Path("pos.md"), output_path=None, destination=Path("pos.docx"))
    assert cli._resolve_paths(args) == (Path("pos.md"), Path("pos.docx"))
    args = argparse.Namespace(input_path=None, source=None, output_path=None, destination=None)
    assert cli._resolve_paths(args) == (None, None)


# --- _format_name ------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("pdf", "pdf"),
        ("PDF", "pdf"),
        ("md", "md"),
        ("markdown", "md"),
        ("docx", "docx"),
        ("report.pdf", "pdf"),
        (Path("x.docx"), "docx"),
        ("  pdf  ", "pdf"),
    ],
)
def test_format_name_recognises_target_formats(value, expected: str) -> None:
    assert cli._format_name(value) == expected


@pytest.mark.parametrize("value", ["", "   ", "xyz", "notes.txt", "dir/pdf", "dir\\pdf", "/tmp/out.docx", Path("a/b.docx")])
def test_format_name_rejects_everything_else(value) -> None:
    assert cli._format_name(value) is None


# --- _coerce_trailing_format ---------------------------------------------------------


def _ns(**kwargs) -> argparse.Namespace:
    defaults = dict(input_path=None, source=None, destination=None, output_path=None, dest_format=None)
    defaults.update(kwargs)
    return argparse.Namespace(**defaults)


def test_coerce_keeps_an_explicit_format() -> None:
    args = _ns(dest_format="md", input_path=Path("a.docx"), source=Path("pdf"))
    cli._coerce_trailing_format(args)
    assert args.dest_format == "md" and args.source == Path("pdf")


def test_coerce_treats_positional_after_input_flag_as_format() -> None:
    args = _ns(input_path=Path("a.docx"), source=Path("pdf"))
    cli._coerce_trailing_format(args)
    assert args.dest_format == "pdf" and args.source is None


def test_coerce_leaves_a_non_format_positional_alone() -> None:
    args = _ns(input_path=Path("a.docx"), source=Path("something_else"))
    cli._coerce_trailing_format(args)
    assert args.dest_format is None and args.source == Path("something_else")


def test_coerce_treats_second_positional_as_format() -> None:
    args = _ns(source=Path("a.md"), destination=Path("docx"))
    cli._coerce_trailing_format(args)
    assert args.dest_format == "docx" and args.destination is None


def test_coerce_keeps_a_real_destination_path() -> None:
    args = _ns(source=Path("a.md"), destination=Path("out/b.docx"))
    cli._coerce_trailing_format(args)
    assert args.dest_format is None and args.destination == Path("out/b.docx")


def test_coerce_ignores_second_positional_when_out_is_given() -> None:
    args = _ns(source=Path("a.md"), destination=Path("pdf"), output_path=Path("o.pdf"))
    cli._coerce_trailing_format(args)
    assert args.dest_format is None


# --- _run_list ---------------------------------------------------------------------


def test_run_list(capsys) -> None:
    assert cli._run_list() == 0
    assert capsys.readouterr().out.splitlines() == [
        "Supported conversions:",
        "  docx -> md",
        "  docx -> pdf",
        "  md -> docx",
        "  pdf -> docx",
    ]


# --- _run_convert (real python-docx / markdown) -----------------------------------------


def _convert(argv: list[str]) -> int:
    return cli.main(argv)


def test_convert_markdown_to_word_with_default_destination(tmp_path, capsys) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# Title\n\n- item", encoding="utf-8")

    assert _convert(["-i", str(source)]) == 0

    destination = tmp_path / "notes.docx"
    assert destination.exists()
    assert [p.text for p in Document(destination).paragraphs] == ["Title", "", "item"]
    assert capsys.readouterr().out == f"Converted {source.resolve()} -> {destination.resolve()}\n"


def test_convert_with_explicit_out_path(tmp_path) -> None:
    source = tmp_path / "a.md"
    source.write_text("# T", encoding="utf-8")
    out = tmp_path / "nested" / "b.docx"
    assert _convert(["-i", str(source), "--out", str(out)]) == 0
    assert out.exists()


def test_convert_with_format_flag_and_trailing_format(tmp_path) -> None:
    doc = tmp_path / "a.docx"
    document = Document()
    document.add_heading("Hello", level=1)
    document.save(doc)

    assert _convert(["-i", str(doc), "-f", "md"]) == 0
    assert (tmp_path / "a.md").read_text(encoding="utf-8") == "# Hello"

    (tmp_path / "a.md").unlink()
    assert _convert(["-i", str(doc), "md"]) == 0
    assert (tmp_path / "a.md").exists()


def test_convert_with_positional_source_and_destination(tmp_path) -> None:
    source, destination = tmp_path / "a.md", tmp_path / "b.docx"
    source.write_text("text", encoding="utf-8")
    assert _convert([str(source), str(destination)]) == 0
    assert destination.exists()


def test_convert_with_positional_source_and_trailing_format(tmp_path) -> None:
    source = tmp_path / "a.md"
    source.write_text("text", encoding="utf-8")
    assert _convert([str(source), "docx"]) == 0
    assert (tmp_path / "a.docx").exists()


def test_convert_subcommand_alias(tmp_path) -> None:
    source = tmp_path / "a.md"
    source.write_text("text", encoding="utf-8")
    assert _convert(["convert", "-i", str(source)]) == 0
    assert (tmp_path / "a.docx").exists()


def test_convert_requires_an_input_file(capsys) -> None:
    assert cli._run_convert(_ns(dest_format="pdf")) == 1
    assert "input file is required" in capsys.readouterr().err


def test_convert_reports_a_missing_input_file(tmp_path, capsys) -> None:
    assert _convert(["-i", str(tmp_path / "missing.md")]) == 1
    assert "input file not found" in capsys.readouterr().err


def test_convert_rejects_same_input_and_output(tmp_path, capsys) -> None:
    source = tmp_path / "a.md"
    source.write_text("x", encoding="utf-8")
    assert _convert(["-i", str(source), "--out", str(source)]) == 1
    assert "must be different files" in capsys.readouterr().err


def test_convert_reports_unsupported_source_extension(tmp_path, capsys) -> None:
    source = tmp_path / "a.txt"
    source.write_text("x", encoding="utf-8")
    assert _convert(["-i", str(source)]) == 1
    assert "No supported conversions for .txt files" in capsys.readouterr().err


def test_convert_reports_unsupported_pair(tmp_path, capsys) -> None:
    source = tmp_path / "a.md"
    source.write_text("x", encoding="utf-8")
    assert _convert(["-i", str(source), "--out", str(tmp_path / "b.pdf")]) == 1
    assert "Unsupported conversion: md -> pdf" in capsys.readouterr().err


def test_convert_reports_import_errors_from_the_converter(monkeypatch, tmp_path, capsys) -> None:
    class Broken:
        def convert(self, source, destination) -> None:
            raise ImportError("Missing required library. Install with: pip install python-docx")

    monkeypatch.setattr(cli, "get_converter", lambda s, d: Broken())
    source = tmp_path / "a.md"
    source.write_text("x", encoding="utf-8")
    assert _convert(["-i", str(source)]) == 1
    assert "Missing required library" in capsys.readouterr().err


def test_convert_reports_unexpected_errors(monkeypatch, tmp_path, capsys) -> None:
    class Failing:
        def convert(self, source, destination) -> None:
            raise RuntimeError("disk on fire")

    monkeypatch.setattr(cli, "get_converter", lambda s, d: Failing())
    source = tmp_path / "a.md"
    source.write_text("x", encoding="utf-8")
    assert _convert(["-i", str(source)]) == 1
    assert "conversion failed: disk on fire" in capsys.readouterr().err


# --- main ----------------------------------------------------------------------------


def test_main_without_arguments_prints_help(capsys) -> None:
    assert cli.main([]) == 0
    assert "usage: pycrux.cli" in capsys.readouterr().out


def test_main_list_and_list_subcommand(capsys) -> None:
    assert cli.main(["--list"]) == 0
    first = capsys.readouterr().out
    assert cli.main(["list"]) == 0
    assert capsys.readouterr().out == first
    assert "md -> docx" in first


def test_main_reads_sys_argv_when_called_without_arguments(monkeypatch, capsys) -> None:
    monkeypatch.setattr(cli.sys, "argv", ["pycrux-convert", "--list"])
    assert cli.main() == 0
    assert "Supported conversions:" in capsys.readouterr().out


def test_main_prints_help_when_nothing_to_do(monkeypatch, capsys) -> None:
    monkeypatch.setattr(cli, "_resolve_paths", lambda args: (None, None))
    assert cli.main(["--"]) == 0
    assert "usage: pycrux.cli" in capsys.readouterr().out


def test_main_rejects_unknown_options() -> None:
    with pytest.raises(SystemExit) as info:
        cli.main(["--bogus"])
    assert info.value.code == 2


# --- root shim -------------------------------------------------------------------------


def test_root_shim_delegates_to_the_package_cli() -> None:
    shim_path = Path(__file__).resolve().parents[1] / "cli.py"
    spec = importlib.util.spec_from_file_location("pycrux_root_cli_shim", shim_path)
    assert spec is not None and spec.loader is not None
    shim = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(shim)
    assert shim.main is cli.main
